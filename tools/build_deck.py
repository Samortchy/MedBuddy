"""
MedBuddy — final Mobile Programming discussion deck (polished).
Team04 — Spring 2026.  Run with the medbuddy-backend conda python.
Output: D:\projects-last-semester\MedBuddy\MedBuddy_Final_Discussion.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

# ── Palette ──────────────────────────────────────────────────────────────────
TEAL      = RGBColor(0x0D, 0x94, 0x88)
TEAL_DK   = RGBColor(0x0A, 0x5C, 0x56)
TEAL_XDK  = RGBColor(0x06, 0x32, 0x2E)
TEAL_SOFT = RGBColor(0xCC, 0xFB, 0xF1)
MINT      = RGBColor(0x5E, 0xEA, 0xD4)
BG_TINT   = RGBColor(0xF2, 0xFB, 0xFA)
INK       = RGBColor(0x0F, 0x17, 0x2A)
GREY      = RGBColor(0x5B, 0x66, 0x70)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
AMBER     = RGBColor(0xF5, 0x9E, 0x0B)
CORAL     = RGBColor(0xFB, 0x71, 0x85)
FONT  = "Segoe UI"
FONTL = "Segoe UI Light"
FONTSB= "Segoe UI Semibold"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height
A = "http://schemas.openxmlformats.org/drawingml/2006/main"

# ── Low-level helpers ────────────────────────────────────────────────────────
def _set(run, size, color=INK, bold=False, italic=False, font=FONT, spacing=None):
    run.font.size = Pt(size); run.font.color.rgb = color
    run.font.bold = bold; run.font.italic = italic; run.font.name = font

def noline(sp): sp.line.fill.background()

def shadow(sp, blur=14, dist=5, alpha=24, col="063A35"):
    spPr = sp._element.spPr
    old = spPr.find(qn('a:effectLst'))
    if old is not None: spPr.remove(old)
    xml = (f'<a:effectLst xmlns:a="{A}">'
           f'<a:outerShdw blurRad="{int(Pt(blur))}" dist="{int(Pt(dist))}" '
           f'dir="5400000" rotWithShape="0">'
           f'<a:srgbClr val="{col}"><a:alpha val="{alpha*1000}"/></a:srgbClr>'
           f'</a:outerShdw></a:effectLst>')
    spPr.append(parse_xml(xml))

def alpha_fill(sp, rgb, pct):
    sp.fill.solid(); sp.fill.fore_color.rgb = rgb
    srgb = sp._element.spPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
    srgb.append(parse_xml(f'<a:alpha xmlns:a="{A}" val="{int(pct*1000)}"/>'))

def grad(target, c1, c2, angle=45):
    """Apply a 2-stop linear gradient to a slide background or shape fill."""
    f = target
    f.gradient()
    stops = f.gradient_stops
    stops[0].position = 0.0; stops[0].color.rgb = c1
    stops[1].position = 1.0; stops[1].color.rgb = c2
    try: f.gradient_angle = angle
    except Exception: pass

def grad_bg(slide, c1, c2, angle=45):
    grad(slide.background.fill, c1, c2, angle)

def rrect(slide, x, y, w, h, color, rad=0.08):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color; noline(sp); sp.shadow.inherit=False
    try: sp.adjustments[0] = rad
    except Exception: pass
    return sp

def rect(slide, x, y, w, h, color):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color; noline(sp); sp.shadow.inherit=False
    return sp

def oval(slide, x, y, w, h, color):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color; noline(sp); sp.shadow.inherit=False
    return sp

def tb(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    t = slide.shapes.add_textbox(x, y, w, h).text_frame
    t.word_wrap = True; t.vertical_anchor = anchor
    return t

def notes(slide, text): slide.notes_slide.notes_text_frame.text = text

# ── Slide chrome (content slides) ────────────────────────────────────────────
def base(slide):
    grad_bg(slide, WHITE, BG_TINT, 60)

def footer(slide, page):
    rect(slide, Inches(0.7), Inches(7.02), Inches(11.93), Pt(1.2), TEAL_SOFT)
    t = tb(slide, Inches(0.7), Inches(7.06), Inches(6), Inches(0.35))
    r = t.paragraphs[0].add_run(); r.text = "MedBuddy  ·  Team04"
    _set(r, 9, GREY, bold=True)
    t = tb(slide, Inches(11.0), Inches(7.06), Inches(1.6), Inches(0.35))
    p = t.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); r.text = f"{page:02d}"; _set(r, 9, TEAL, bold=True)

def watermark(slide, n):
    t = tb(slide, Inches(10.4), Inches(0.05), Inches(2.9), Inches(1.8))
    p = t.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); r.text = f"{n:02d}"
    _set(r, 80, TEAL_SOFT, bold=True, font=FONTL)

def header(slide, title, kicker, n, page):
    base(slide); watermark(slide, n)
    rrect(slide, Inches(0.7), Inches(0.62), Inches(0.16), Inches(0.62), TEAL, rad=0.5)
    t = tb(slide, Inches(1.0), Inches(0.5), Inches(9.2), Inches(1.2))
    p = t.paragraphs[0]
    r = p.add_run(); r.text = kicker.upper(); _set(r, 12, TEAL, bold=True)
    p2 = t.add_paragraph()
    r = p2.add_run(); r.text = title; _set(r, 31, INK, bold=True)
    footer(slide, page)

# Accent-bar feature card
def feat_card(slide, x, y, w, h, title, desc, accent=TEAL, tint=WHITE):
    card = rrect(slide, x, y, w, h, tint, rad=0.10); shadow(card, blur=12, alpha=16)
    rrect(slide, x, y, Inches(0.12), h, accent, rad=0.5)
    t = tb(slide, x+Inches(0.32), y+Inches(0.1), w-Inches(0.5), h-Inches(0.2),
           anchor=MSO_ANCHOR.MIDDLE)
    p = t.paragraphs[0]; r = p.add_run(); r.text = title; _set(r, 14.5, INK, bold=True)
    p2 = t.add_paragraph(); p2.space_before = Pt(2)
    r = p2.add_run(); r.text = desc; _set(r, 10.5, GREY)

def bullets(slide, items, x=Inches(1.0), y=Inches(2.0),
            w=Inches(11.3), h=Inches(4.7), size=17, gap=12):
    t = tb(slide, x, y, w, h); first=True
    for it in items:
        lvl = 0
        if isinstance(it, tuple): it, lvl = it
        p = t.paragraphs[0] if first else t.add_paragraph(); first=False
        p.space_after = Pt(gap); p.level = lvl
        dot = p.add_run(); dot.text = ("●  " if lvl==0 else "–  ")
        _set(dot, (size-5) if lvl==0 else (size-6), TEAL if lvl==0 else CORAL, bold=True)
        r = p.add_run(); r.text = it
        _set(r, size if lvl==0 else size-2, INK if lvl==0 else GREY)
    return t

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — COVER
# ════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
grad_bg(s, TEAL_XDK, TEAL_DK, 45)
# decorative blobs
b = oval(s, Inches(8.6), Inches(-1.8), Inches(6.5), Inches(6.5), TEAL); alpha_fill(b, TEAL, 28)
b = oval(s, Inches(10.6), Inches(3.0), Inches(5.0), Inches(5.0), MINT); alpha_fill(b, MINT, 16)
b = oval(s, Inches(-1.6), Inches(4.6), Inches(4.2), Inches(4.2), AMBER); alpha_fill(b, AMBER, 12)
# top kicker
t = tb(s, Inches(0.95), Inches(0.7), Inches(11), Inches(0.5))
r = t.paragraphs[0].add_run()
r.text = "MOBILE PROGRAMMING   ·   FINAL DISCUSSION   ·   SPRING 2026"
_set(r, 13, MINT, bold=True)
# pulse mark + title
rrect(s, Inches(0.95), Inches(2.05), Inches(0.55), Inches(0.55), AMBER, rad=0.3)
t = tb(s, Inches(0.9), Inches(2.55), Inches(11.6), Inches(2.6))
p = t.paragraphs[0]
r = p.add_run(); r.text = "MedBuddy"; _set(r, 72, WHITE, bold=True)
p2 = t.add_paragraph()
r = p2.add_run(); r.text = "AI-Powered Elderly Health Companion"
_set(r, 26, TEAL_SOFT, font=FONTL)
# accent line
rrect(s, Inches(0.98), Inches(5.05), Inches(2.2), Pt(4), AMBER, rad=0.5)
# team card
card = rrect(s, Inches(0.95), Inches(5.35), Inches(7.9), Inches(1.65), TEAL_DK, rad=0.12)
alpha_fill(card, WHITE, 10)
t = tb(s, Inches(1.25), Inches(5.5), Inches(7.4), Inches(1.4), anchor=MSO_ANCHOR.MIDDLE)
p = t.paragraphs[0]; r = p.add_run(); r.text = "Team04"
_set(r, 16, AMBER, bold=True)
team = "Ahmed Samer  ·  Ali Abdallah  ·  Ismail Hesham  ·  Abu Bakr Hegazy  ·  Mena Khaled"
p2 = t.add_paragraph(); r = p2.add_run(); r.text = team; _set(r, 12.5, WHITE)
ids = "2022/08211 · 2022/05974 · 2022/00106 · 2022/02645 · 2022/03469"
p3 = t.add_paragraph(); r = p3.add_run(); r.text = ids; _set(r, 10.5, TEAL_SOFT)
notes(s, "MedBuddy is an AI-powered companion that helps elderly patients manage "
         "medications, daily wellness, and emergencies, while giving caregivers remote "
         "visibility. Built with Flutter, FastAPI, Supabase and several AI services. "
         "Team04 — all members contributed across the stack.")

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — PROBLEM
# ════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "The Problem We Solve", "Motivation", 1, 2)
bullets(s, [
    "Elderly patients often forget medications or take the wrong dose.",
    "Subtle declines in mood, sleep, energy and pain go unnoticed until serious.",
    "In a fall or emergency, every minute matters — but help is often far away.",
    "Family caregivers want reassurance but cannot watch over a parent 24/7.",
    "Most health apps are built for tech-savvy users, not seniors.",
], size=18, gap=15, y=Inches(2.05), h=Inches(3.6))
band = rrect(s, Inches(1.0), Inches(5.55), Inches(11.3), Inches(1.05), TEAL, rad=0.14)
shadow(band, blur=14, alpha=22)
t = tb(s, Inches(1.4), Inches(5.6), Inches(10.6), Inches(0.95), anchor=MSO_ANCHOR.MIDDLE)
r = t.paragraphs[0].add_run()
r.text = ("Our goal:  one simple app that keeps seniors safe and healthy — and keeps "
          "their caregivers connected and informed.")
_set(r, 15.5, WHITE, bold=True, italic=True)
notes(s, "Frame the real-world need. Two user types: the patient (senior) and the "
         "caregiver (family member). MedBuddy bridges them.")

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — TWO ROLES
# ════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "Two Roles, One App", "Solution Overview", 2, 3)
def role_card(x, head, items, accent):
    card = rrect(s, x, Inches(2.0), Inches(5.65), Inches(4.65), WHITE, rad=0.07)
    shadow(card, blur=16, alpha=18)
    rrect(s, x, Inches(2.0), Inches(5.65), Inches(0.85), accent, rad=0.07)
    rect(s, x, Inches(2.45), Inches(5.65), Inches(0.4), accent)  # square off bottom of band
    t = tb(s, x+Inches(0.35), Inches(2.05), Inches(5), Inches(0.8), anchor=MSO_ANCHOR.MIDDLE)
    r = t.paragraphs[0].add_run(); r.text = head; _set(r, 21, WHITE, bold=True)
    t = tb(s, x+Inches(0.4), Inches(3.05), Inches(4.9), Inches(3.45))
    first=True
    for it in items:
        p = t.paragraphs[0] if first else t.add_paragraph(); first=False
        p.space_after = Pt(8)
        d = p.add_run(); d.text="●  "; _set(d, 11, accent, bold=True)
        r = p.add_run(); r.text = it; _set(r, 14, INK)
role_card(Inches(0.75), "Patient", [
    "Medication schedule & reminders", "Daily wellness check-ins",
    "AI health buddy (chat + voice)", "Symptom logging with AI triage",
    "One-tap SOS on every screen", "Personal health history"], TEAL)
role_card(Inches(6.9), "Caregiver", [
    "Link to a patient via secure code", "Read-only health dashboards",
    "Wellness, meds, symptoms, emergencies", "Two-way secure messaging",
    "Instant emergency push alerts", "Auto-join emergency video call"], TEAL_DK)
notes(s, "Role lives in the Supabase JWT user_metadata and drives the entire UI. A "
         "patient and caregiver are connected through a link created by an invite code.")

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — ARCHITECTURE
# ════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "System Architecture", "How It Fits Together", 3, 4)
def node(x, y, w, h, title, sub, c1, c2):
    sp = rrect(s, x, y, w, h, c1, rad=0.12); grad(sp.fill, c1, c2, 60); shadow(sp, blur=14, alpha=22)
    t = sp.text_frame; t.word_wrap=True; t.vertical_anchor=MSO_ANCHOR.MIDDLE
    p = t.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    r = p.add_run(); r.text=title; _set(r, 16, WHITE, bold=True)
    p2 = t.add_paragraph(); p2.alignment=PP_ALIGN.CENTER
    r = p2.add_run(); r.text=sub; _set(r, 10.5, RGBColor(0xE6,0xFF,0xFA))
    return sp
def conn(x, y, w):
    a = s.shapes.add_shape(MSO_SHAPE.LEFT_RIGHT_ARROW, x, y, w, Inches(0.3))
    a.fill.solid(); a.fill.fore_color.rgb=AMBER; noline(a); a.shadow.inherit=False
yN = Inches(2.55)
node(Inches(0.75), yN, Inches(3.25), Inches(1.55), "Flutter App",
     "Patient + Caregiver UI\nRiverpod · Dio", TEAL, MINT)
node(Inches(5.05), yN, Inches(3.25), Inches(1.55), "FastAPI Backend",
     "REST API · workers\nservice-role access", TEAL_DK, TEAL)
node(Inches(9.35), yN, Inches(3.25), Inches(1.55), "Supabase",
     "Postgres + Auth (JWT)", RGBColor(0x2F,0x9E,0x44), RGBColor(0x69,0xC9,0x7B))
conn(Inches(4.0), Inches(3.05), Inches(1.1))
conn(Inches(8.3), Inches(3.05), Inches(1.1))
# AI services
ai = node(Inches(5.05), Inches(4.75), Inches(3.25), Inches(1.45), "AI Services",
          "LLM · STT · TTS\nAgora · FCM", AMBER, RGBColor(0xFB,0xBF,0x24))
ai.text_frame.paragraphs[0].runs[0].font.color.rgb = INK
for p in ai.text_frame.paragraphs:
    for r in p.runs: r.font.color.rgb = INK
rect(s, Inches(6.6), Inches(4.1), Pt(3), Inches(0.65), AMBER)
t = tb(s, Inches(0.95), Inches(6.45), Inches(11.4), Inches(0.55))
r = t.paragraphs[0].add_run()
r.text = ("Auth-only Supabase:  Flutter authenticates with Supabase; ALL data flows "
          "through FastAPI, which holds the service-role key.")
_set(r, 13.5, GREY, italic=True)
notes(s, "Key decision: the Flutter app never touches the DB directly. It authenticates "
         "with Supabase, gets a JWT, sends it as a Bearer token to FastAPI, which "
         "validates it and uses the service-role key for all DB access. Centralizes "
         "logic and security on the backend.")

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — TECH STACK
# ════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "Technology Stack", "What We Built With", 4, 5)
cols = [
    ("Frontend", TEAL, ["Flutter 3.x / Dart 3.11", "Riverpod 2.5 (state)",
        "Dio 5.4 (HTTP)", "sqflite (local cache)", "image_picker", "permission_handler"]),
    ("Backend", TEAL_DK, ["FastAPI (Python)", "supabase-py (service role)",
        "Background workers", "slowapi (rate limit)", "Pydantic models", "REST + JSON"]),
    ("Data & Auth", RGBColor(0x2F,0x9E,0x44), ["Supabase Postgres", "Supabase Auth (JWT)",
        "Role in user_metadata", "Row-based ownership", "Soft-delete", "Unique indexes"]),
    ("AI & Realtime", AMBER, ["OpenRouter · Llama 3.3", "faster-whisper (STT)",
        "Chatterbox (TTS)", "Agora RTC (video)", "Firebase FCM (push)", "On-GPU inference"]),
]
x = Inches(0.7); cw = Inches(3.0); g = Inches(0.13)
for title, accent, items in cols:
    card = rrect(s, x, Inches(2.0), cw, Inches(4.75), WHITE, rad=0.07)
    shadow(card, blur=14, alpha=16)
    rrect(s, x, Inches(2.0), cw, Inches(0.7), accent, rad=0.07)
    rect(s, x, Inches(2.4), cw, Inches(0.3), accent)
    hd = tb(s, x, Inches(2.05), cw, Inches(0.6), anchor=MSO_ANCHOR.MIDDLE)
    p = hd.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    fg = INK if accent==AMBER else WHITE
    r = p.add_run(); r.text=title; _set(r, 15, fg, bold=True)
    t = tb(s, x+Inches(0.22), Inches(2.85), cw-Inches(0.4), Inches(3.8))
    first=True
    for it in items:
        p = t.paragraphs[0] if first else t.add_paragraph(); first=False
        p.space_after = Pt(7)
        d=p.add_run(); d.text="• "; _set(d,12,accent,bold=True)
        r=p.add_run(); r.text=it; _set(r,11.5,INK)
    x = x + cw + g
notes(s, "The 'technologies' rubric item. Justify each: Riverpod for testable reactive "
         "state, Dio for interceptors, Supabase for auth+managed Postgres, OpenRouter to "
         "access Llama cheaply, Agora for low-latency video, FCM for reliable push.")

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — PATIENT FEATURES
# ════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "Patient Features", "Core Functionality", 5, 6)
feats = [
    ("Guided onboarding", "Profile, conditions, meds, contacts & prefs."),
    ("Medication schedule", "Reminders, mark-as-taken, adherence history."),
    ("Wellness check-in", "Mood, energy, sleep, pain — trend charts."),
    ("AI health buddy", "Context-aware chat + voice (STT / TTS)."),
    ("Symptom log", "Free-text symptoms with AI severity triage."),
    ("History hub", "Wellness, symptoms, meds, emergencies."),
    ("SOS button", "One tap on every screen → emergency flow."),
    ("Profile photo", "Camera/gallery upload, cached locally."),
]
accents = [TEAL, TEAL_DK, TEAL, AMBER, CORAL, TEAL_DK, CORAL, TEAL]
x0,y0 = Inches(0.75), Inches(2.0)
for i,(t_,d_) in enumerate(feats):
    col=i%2; row=i//2
    feat_card(s, x0+col*Inches(6.0), y0+row*Inches(1.2), Inches(5.75), Inches(1.05),
              t_, d_, accent=accents[i])
notes(s, "Patient journey: onboard → daily reminders & check-ins → talk to the AI buddy "
         "→ log symptoms → SOS in trouble. Emphasize senior-friendly UI: big text, high "
         "contrast, simple navigation.")

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 7 — CAREGIVER FEATURES
# ════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "Caregiver Features", "Remote Care", 6, 7)
feats = [
    ("Secure linking", "6-digit invite code links caregiver ↔ patient."),
    ("Patient dashboards", "Read-only views of all patient data."),
    ("Wellness & symptoms", "Check-in trends and AI-flagged symptoms."),
    ("Medication view", "Patient's active meds and adherence."),
    ("Emergency log", "History of SOS / fall events & outcomes."),
    ("Secure messaging", "Two-way chat thread with the patient."),
    ("Push alerts (FCM)", "Instant notice when an emergency fires."),
    ("Auto-join call", "Ringtone + auto-join the emergency call."),
]
accents = [TEAL_DK, TEAL, TEAL, TEAL_DK, CORAL, TEAL, AMBER, CORAL]
for i,(t_,d_) in enumerate(feats):
    col=i%2; row=i//2
    feat_card(s, x0+col*Inches(6.0), y0+row*Inches(1.2), Inches(5.75), Inches(1.05),
              t_, d_, accent=accents[i])
notes(s, "The caregiver side turns scattered data into peace of mind. All access is "
         "gated: a caregiver only sees patients they are explicitly linked to, verified "
         "on every backend request.")

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 8 — AI INTEGRATION
# ════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "AI Integration", "The Smart Layer", 7, 8)
bullets(s, [
    "Context-aware chat — the backend injects the patient's name, age, language, "
    "conditions and medications into the system prompt, so the AI answers personally.",
    ("Server-side prompt building keeps patient data secure and the client thin.", 1),
    "Voice — faster-whisper turns speech into text; Chatterbox turns the reply into "
    "natural speech for hands-free use.",
    "Symptom triage — each symptom is classified normal / watch / flagged; an "
    "emergency keyword list short-circuits to an instant flag.",
    ("Flagged symptoms automatically notify linked caregivers.", 1),
    "Visit summaries — an LLM turns a doctor-visit transcript into structured notes "
    "(diagnosis, medications, follow-ups).",
    "Responses stay short, warm, and in the patient's preferred language.",
], size=15.5, gap=10, y=Inches(2.0), h=Inches(4.8))
notes(s, "AI is our differentiator. Q&A point: the LLM does NOT get patient data from the "
         "client — FastAPI fetches it from the DB and builds the prompt. Model: Llama 3.3 "
         "70B via OpenRouter. STT/TTS run locally on GPU. Triage mixes a deterministic "
         "keyword check with the LLM.")

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 9 — EMERGENCY FLOW
# ════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "Emergency & SOS Flow", "When It Matters Most", 8, 9)
steps = [
    ("1","Trigger","Patient taps SOS (or a fall is detected) on any screen."),
    ("2","Verify","A quick liveness check filters out false alarms."),
    ("3","Notify","Backend pushes FCM alerts to all linked caregivers."),
    ("4","Connect","Agora video call opens; caregiver phone rings & auto-joins."),
    ("5","Escalate","A background worker escalates if no caregiver responds."),
]
y = Inches(2.2)
# vertical connector
rect(s, Inches(1.22), Inches(2.45), Pt(2.5), Inches(3.7), TEAL_SOFT)
for num,t_,d_ in steps:
    c = oval(s, Inches(0.9), y, Inches(0.66), Inches(0.66), TEAL)
    grad(c.fill, TEAL, TEAL_DK, 60); shadow(c, blur=10, alpha=26)
    tf = c.text_frame; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    r=p.add_run(); r.text=num; _set(r,18,WHITE,bold=True)
    card = rrect(s, Inches(1.85), y-Inches(0.06), Inches(10.5), Inches(0.78), WHITE, rad=0.14)
    shadow(card, blur=10, alpha=14)
    tx = tb(s, Inches(2.1), y-Inches(0.04), Inches(10.1), Inches(0.74), anchor=MSO_ANCHOR.MIDDLE)
    p = tx.paragraphs[0]
    r=p.add_run(); r.text=t_+"    "; _set(r,16,TEAL_DK,bold=True)
    r=p.add_run(); r.text=d_; _set(r,14,INK)
    y = y + Inches(0.82)
notes(s, "Highest-stakes feature. SOS is reachable from EVERY patient screen via a "
         "persistent overlay button. Combines FCM (alerting) and Agora RTC (live video) "
         "with a worker-based escalation safety net.")

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 10 — STATE & NAVIGATION
# ════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "State Management & Navigation", "Architecture in the App", 9, 10)
bullets(s, [
    "Riverpod 2.5 — StateNotifierProvider for mutable flows (meds, symptoms, chat) and "
    "FutureProvider.family for async, per-patient reads.",
    "Single Dio instance with an interceptor that injects the Bearer token on every "
    "request and signs out on 401.",
    ("Configurable base URL via --dart-define=API_BASE for emulator vs real device.", 1),
    "Named routes drive navigation; the user's role selects the patient vs caregiver "
    "shell and bottom-nav destinations.",
    "Providers are invalidated after edits, so screens refresh with fresh server data.",
    "Defensive patterns — display de-duplication, null-guards with loading/retry states "
    "instead of fake placeholder data.",
], size=16.5, gap=12, y=Inches(2.05), h=Inches(4.7))
notes(s, "Covers the state-management and navigation rubric. Riverpod = compile-safe, "
         "testable, reactive state. Dio interceptor centralizes auth. Explain provider "
         "invalidation after an edit screen pops.")

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 11 — DEVICE FEATURES & PERSISTENCE
# ════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "Device Features & Persistence", "Native Capabilities", 10, 11)
def list_card(x, head, accent, items):
    card = rrect(s, x, Inches(2.0), Inches(5.75), Inches(4.65), WHITE, rad=0.07)
    shadow(card, blur=16, alpha=16)
    rrect(s, x, Inches(2.0), Inches(5.75), Inches(0.7), accent, rad=0.07)
    rect(s, x, Inches(2.4), Inches(5.75), Inches(0.3), accent)
    hd = tb(s, x+Inches(0.35), Inches(2.05), Inches(5), Inches(0.6), anchor=MSO_ANCHOR.MIDDLE)
    r = hd.paragraphs[0].add_run(); r.text=head
    _set(r, 17, INK if accent==AMBER else WHITE, bold=True)
    t = tb(s, x+Inches(0.4), Inches(2.95), Inches(5), Inches(3.5))
    first=True
    for it in items:
        p = t.paragraphs[0] if first else t.add_paragraph(); first=False
        p.space_after = Pt(9)
        d=p.add_run(); d.text="●  "; _set(d,11,accent,bold=True)
        r=p.add_run(); r.text=it; _set(r,13.5,INK)
list_card(Inches(0.75), "Device Features", TEAL, [
    "Camera & Gallery — profile photo (image_picker)",
    "Microphone & Audio — STT / TTS / voice",
    "Push notifications — Firebase FCM",
    "Video & audio call — Agora RTC",
    "Runtime permissions — permission_handler",
    "Ringtone on incoming emergency"])
list_card(Inches(6.85), "Data Persistence", AMBER, [
    "sqflite local database on device",
    "profile_images table — full CRUD",
    "Photos cached locally per user",
    "Supabase Postgres = source of truth",
    "Soft-delete + unique indexes server-side",
    "Survives app restarts"])
notes(s, "Maps to rubric items: device features (camera, mic, push, video) and data "
         "persistence (sqflite). Profile images stored locally with full CRUD; the "
         "server DB stays the source of truth for health data.")

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 12 — INTERACTIVITY & UX
# ════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "Interactivity & UX Polish", "Senior-Friendly Design", 11, 12)
bullets(s, [
    "Swipe-to-delete with Undo — symptom entries use a Dismissible plus a floating "
    "Undo SnackBar to prevent accidental loss.",
    "Consistent theming — one teal design system (colors, dimens, text styles) across "
    "every screen.",
    "Accessibility — large tap targets, high contrast, clamped text scaling so layouts "
    "never break when the OS font is enlarged.",
    "Portrait orientation lock for a predictable senior experience.",
    "Responsive layouts — Wrap / Flexible widgets adapt to long content and screen sizes.",
    "Clear feedback — loading spinners, retry states, and success / error messages.",
], size=16.5, gap=12, y=Inches(2.05), h=Inches(4.7))
notes(s, "Covers interactivity and responsiveness/UX. Dismissible + Undo is the concrete "
         "interactivity example. Theming + text scaling + orientation lock = responsive, "
         "consistent UX for elderly users.")

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 13 — CHALLENGES & LEARNINGS
# ════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "Challenges & Learnings", "What We Took Away", 12, 13)
list_card(Inches(0.75), "Challenges", CORAL, [
    "Auth-only Supabase + backend service role",
    "Giving the AI real patient context securely",
    "Running GPU AI services (STT / TTS) locally",
    "Emulator vs real-device networking",
    "Duplicate-data bugs → indexes + guards"])
list_card(Inches(6.85), "What We Learned", TEAL, [
    "End-to-end mobile + backend + AI integration",
    "Secure auth & role-based access design",
    "Reactive state management at scale",
    "Real-time features (push + video)",
    "Shipping a polished, accessible product"])
notes(s, "Be honest and specific about challenges — examiners respect this. Each learning "
         "ties back to a rubric area.")

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 14 — THANK YOU
# ════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
grad_bg(s, TEAL_XDK, TEAL_DK, 45)
b = oval(s, Inches(-2.0), Inches(-2.2), Inches(7.5), Inches(7.5), TEAL); alpha_fill(b, TEAL, 24)
b = oval(s, Inches(9.5), Inches(3.5), Inches(6.0), Inches(6.0), MINT); alpha_fill(b, MINT, 14)
b = oval(s, Inches(10.8), Inches(-1.6), Inches(3.4), Inches(3.4), AMBER); alpha_fill(b, AMBER, 16)
rrect(s, Inches(1.05), Inches(2.45), Inches(0.55), Inches(0.55), AMBER, rad=0.3)
t = tb(s, Inches(1.0), Inches(2.95), Inches(11), Inches(2.0))
p = t.paragraphs[0]; r = p.add_run(); r.text = "Thank You"
_set(r, 76, WHITE, bold=True)
p2 = t.add_paragraph(); r = p2.add_run()
r.text = "We're happy to answer your questions."
_set(r, 24, TEAL_SOFT, font=FONTL)
rrect(s, Inches(1.08), Inches(5.05), Inches(2.2), Pt(4), AMBER, rad=0.5)
t = tb(s, Inches(1.0), Inches(5.35), Inches(11.5), Inches(1.4))
p = t.paragraphs[0]; r = p.add_run(); r.text = "MedBuddy  ·  Team04"
_set(r, 16, AMBER, bold=True)
p2 = t.add_paragraph(); r = p2.add_run()
r.text = "Ahmed Samer  ·  Ali Abdallah  ·  Ismail Hesham  ·  Abu Bakr Hegazy  ·  Mena Khaled"
_set(r, 13, WHITE)
p3 = t.add_paragraph(); r = p3.add_run()
r.text = "Mobile Programming · Final Discussion · Spring 2026"
_set(r, 11, TEAL_SOFT)
notes(s, "Thank the panel, restate the one-line value proposition, and open the floor "
         "for the technical Q&A.")

import os, time
base = r"D:\projects-last-semester\MedBuddy\MedBuddy_Final_Discussion.pptx"
out = base
try:
    prs.save(out)
except PermissionError:
    out = base.replace(".pptx", f"_v{time.strftime('%H%M%S')}.pptx")
    prs.save(out)
n = len(prs.slides._sldIdLst)
print("Saved:", out, "-", n, "slides")
